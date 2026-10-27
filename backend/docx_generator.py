import os
from typing import List, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# 색상 팔레트 (신뢰감 있는 엔터프라이즈 네이비 & 블루)
COLOR_PRIMARY = RGBColor(26, 54, 93)       # #1A365D 딥 네이비
COLOR_SECONDARY = RGBColor(43, 108, 176)   # #2B6CB0 미드 네이비
COLOR_ACCENT = RGBColor(13, 148, 136)      # #0D9488 틸
COLOR_DARK = RGBColor(30, 41, 59)          # #1E293B 슬레이트 다크
COLOR_MUTED = RGBColor(100, 116, 139)      # #64748B 슬레이트 그레이

HEX_PRIMARY = "1A365D"
HEX_LIGHT_BG = "F8FAFC"
HEX_ALT_ROW = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_HIGHLIGHT = "E0F2FE"

def set_cell_background(cell, hex_color: str):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(margin_name)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, hex_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{HEX_PRIMARY}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_font(run, font_name="Malgun Gothic", size_pt=10, bold=False, color=COLOR_DARK):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = color
    # 한글 폰트 적용을 위한 rFonts 설정
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:eastAsia="{font_name}" w:hAnsi="{font_name}"/>')
    rPr.append(rFonts)


def build_docx_report(records: List[Dict[str, Any]], kpis: Dict[str, Any], ai_analysis: Dict[str, Any], output_path: str):
    """
    경영진 보고용 정식 워드 성과보고서(.docx) 문서를 생성합니다.
    """
    doc = Document()

    # 페이지 여백 설정 (상하 25mm, 좌우 25mm)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.98)
        s.bottom_margin = Inches(0.98)
        s.left_margin = Inches(0.98)
        s.right_margin = Inches(0.98)

    # ==========================================
    # 1. 표지 (Cover Page)
    # ==========================================
    cover_space = doc.add_paragraph()
    cover_space.paragraph_format.space_before = Pt(80)

    # 뱃지 라벨
    p_badge = doc.add_paragraph()
    r_badge = p_badge.add_run("OFFICIAL PERFORMANCE REPORT  |  경영진 보고용")
    format_font(r_badge, size_pt=10, bold=True, color=COLOR_SECONDARY)
    p_badge.paragraph_format.space_after = Pt(12)

    # 메인 타이틀
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("트릿지(Tridge) 솔루션 활용\n전사 성과 및 비즈니스 ROI 보고서")
    format_font(r_title, size_pt=24, bold=True, color=COLOR_PRIMARY)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.paragraph_format.line_spacing = 1.2

    # 서브 타이틀
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("부서별 활용 실적 분석, 재무 기여도(원가절감·수출) 집계 및 차년도 활용 전략")
    format_font(r_sub, size_pt=12, bold=False, color=COLOR_MUTED)
    p_sub.paragraph_format.space_after = Pt(180)

    # 메타데이터 표
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_data = [
        ("보고 일자", "2026년 10월 26일"),
        ("대상 기간", "2026년 하반기 (Tridge Eye / Ark / Data Analytics)"),
        ("참여 부서", f"{kpis.get('department_count', 0)}개 부서 ({', '.join(kpis.get('departments', []))})"),
        ("주관 부서", "전략기획팀 / 트릿지 도입 총괄 TF")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.3)
        c1.width = Inches(4.5)
        set_cell_margins(c0, 60, 60, 60, 60)
        set_cell_margins(c1, 60, 60, 60, 60)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        format_font(r0, size_pt=9.5, bold=True, color=COLOR_MUTED)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        format_font(r1, size_pt=9.5, bold=False, color=COLOR_DARK)

    doc.add_page_break()

    # ==========================================
    # 2. 제1장. Executive Summary (경영진 핵심 요약)
    # ==========================================
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("제1장. Executive Summary (경영진 핵심 요약)")
    format_font(r_h1, size_pt=16, bold=True, color=COLOR_PRIMARY)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(12)

    # AI 요약 문단들
    summary_paras = ai_analysis.get("executive_summary_paragraphs", [])
    for p_text in summary_paras:
        p = doc.add_paragraph()
        r = p.add_run(p_text)
        format_font(r, size_pt=10.5, color=COLOR_DARK)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(8)

    # 3대 핵심 성과 카드 (Callout Box)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    p_hl_title = doc.add_paragraph()
    r_hl_t = p_hl_title.add_run("❖ 금번 분기 3대 핵심 성과 (Key Highlights)")
    format_font(r_hl_t, size_pt=11.5, bold=True, color=COLOR_SECONDARY)
    p_hl_title.paragraph_format.space_after = Pt(8)

    highlights = ai_analysis.get("highlights", [])
    hl_table = doc.add_table(rows=len(highlights), cols=1)
    hl_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(hl_table, HEX_BORDER)

    for idx, hl in enumerate(highlights):
        cell = hl_table.rows[idx].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, HEX_LIGHT_BG if idx % 2 == 0 else "FFFFFF")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        p = cell.paragraphs[0]
        r_num = p.add_run(f"[{idx+1}] ")
        format_font(r_num, size_pt=10, bold=True, color=COLOR_ACCENT)
        r_t = p.add_run(hl.get("title", "") + "\n")
        format_font(r_t, size_pt=10, bold=True, color=COLOR_PRIMARY)
        r_d = p.add_run(hl.get("desc", ""))
        format_font(r_d, size_pt=9.5, bold=False, color=COLOR_DARK)
        p.paragraph_format.line_spacing = 1.2

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # ==========================================
    # 3. 제2장. 전사 정량 KPI 총괄 대시보드
    # ==========================================
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("제2장. 전사 정량 KPI 총괄 현황")
    format_font(r_h2, size_pt=16, bold=True, color=COLOR_PRIMARY)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(12)

    # 4개 핵심 지표 카드 테이블
    kpi_table = doc.add_table(rows=2, cols=4)
    kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(kpi_table, HEX_PRIMARY)

    kpi_boxes = [
        ("직접 재무 기여액", f"{kpis.get('total_financial_impact', 0):,.0f} 만원", "원가 절감 및 수출 파이프라인"),
        ("리서치 절감 시간", f"{kpis.get('total_time_saved', 0):,.0f} 시간", "기존 방식 대비 70% 단축"),
        ("타깃 DB 확보", f"{kpis.get('total_db_count', 0)} 개사", "글로벌 바이어/공급사"),
        ("솔루션 만족도", f"{kpis.get('avg_satisfaction', 0)} / 5.0", "전사 평균 만족도")
    ]

    for idx, (label, val, desc) in enumerate(kpi_boxes):
        c_top = kpi_table.rows[0].cells[idx]
        c_bottom = kpi_table.rows[1].cells[idx]
        c_top.width = Inches(1.6)
        c_bottom.width = Inches(1.6)
        
        set_cell_background(c_top, HEX_PRIMARY)
        set_cell_background(c_bottom, HEX_LIGHT_BG)
        set_cell_margins(c_top, 100, 100, 100, 100)
        set_cell_margins(c_bottom, 120, 120, 100, 100)
        
        p0 = c_top.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(label)
        format_font(r0, size_pt=9.5, bold=True, color=RGBColor(255, 255, 255))
        
        p1 = c_bottom.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(val + "\n")
        format_font(r1, size_pt=13, bold=True, color=COLOR_PRIMARY)
        r2 = p1.add_run(desc)
        format_font(r2, size_pt=8, bold=False, color=COLOR_MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 솔루션별 점유 현황 표
    p_sol_title = doc.add_paragraph()
    r_sol_t = p_sol_title.add_run("❖ 솔루션별 및 진행상태별 집계")
    format_font(r_sol_t, size_pt=11.5, bold=True, color=COLOR_SECONDARY)
    p_sol_title.paragraph_format.space_after = Pt(6)

    breakdown_table = doc.add_table(rows=1, cols=4)
    breakdown_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(breakdown_table, HEX_BORDER)

    headers = ["구분", "트릿지 아이 (Eye)", "트릿지 아크 (Ark)", "Data Analytics"]
    for i, h in enumerate(headers):
        cell = breakdown_table.rows[0].cells[i]
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, 120, 120, 120, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        format_font(r, size_pt=9.5, bold=True, color=RGBColor(255, 255, 255))

    sol_map = kpis.get("solution_breakdown", {})
    row_data = [
        ("활용 건수", f"{sol_map.get('트릿지 아이', 0)} 건", f"{sol_map.get('트릿지 아크', 0)} 건", f"{sol_map.get('Data Analytics', 0)} 건"),
        ("주요 활용 분야", "시세/가격 추이, 공급처 발굴", "글로벌 프로모션, 바이어 매칭", "시장 트렌드 분석, PB 기획")
    ]
    for r_idx, r_vals in enumerate(row_data):
        row = breakdown_table.add_row()
        for c_idx, val in enumerate(r_vals):
            cell = row.cells[c_idx]
            set_cell_background(cell, HEX_ALT_ROW if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 100, 100, 120, 120)
            p = cell.paragraphs[0]
            if c_idx > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            format_font(r, size_pt=9, bold=(c_idx == 0), color=COLOR_DARK)

    doc.add_page_break()

    # ==========================================
    # 4. 제3장. 부서별 세부 성과 분석 (Deep Dive)
    # ==========================================
    h3 = doc.add_paragraph()
    r_h3 = h3.add_run("제3장. 부서별 세부 활용 실적 및 비즈니스 성과 (Deep Dive)")
    format_font(r_h3, size_pt=16, bold=True, color=COLOR_PRIMARY)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(12)

    for idx, rec in enumerate(records):
        dept = rec.get("department", "기타 부서")
        author = rec.get("author", "-")
        sol = rec.get("solution", "-")
        status = rec.get("status", "-")
        period = rec.get("period", "-")
        goal = rec.get("goal", "-")
        target = rec.get("target", "-")
        db_cnt = rec.get("valid_db_count", 0)
        action_cnt = rec.get("action_count", 0)
        fin = rec.get("financial_impact", 0)
        t_saved = rec.get("time_saved_hours", 0)
        insight = rec.get("qualitative_insight", "-")
        next_p = rec.get("next_plan", "-")
        sat = rec.get("satisfaction", 4)

        # 부서 헤더
        p_dept = doc.add_paragraph()
        r_num = p_dept.add_run(f"3.{idx+1} ")
        format_font(r_num, size_pt=13, bold=True, color=COLOR_SECONDARY)
        r_dept = p_dept.add_run(f"{dept} 파트 실적 분석 ")
        format_font(r_dept, size_pt=13, bold=True, color=COLOR_PRIMARY)
        r_badge = p_dept.add_run(f"[{sol} / 상태: {status} / 만족도: {'★'*sat}{'☆'*(5-sat)}]")
        format_font(r_badge, size_pt=9.5, bold=False, color=COLOR_MUTED)
        p_dept.paragraph_format.space_before = Pt(12)
        p_dept.paragraph_format.space_after = Pt(6)

        # 세부 테이블
        d_table = doc.add_table(rows=6, cols=2)
        d_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(d_table, HEX_BORDER)

        items = [
            ("비즈니스 목표", goal),
            ("타깃 품목 및 국가", target),
            ("정량적 비즈니스 성과", f"• 유효 DB 확보: {db_cnt}개사   • 직접 액션: {action_cnt}건\n• 재무 기여액: {fin:,.0f}만 원   • 리서치 절감: {t_saved:,.0f}시간"),
            ("핵심 성과 및 시장 인사이트", insight),
            ("세부 활용 내용", rec.get("activity", "-")),
            ("차기 계획 및 목표", next_p)
        ]

        for row_i, (k, v) in enumerate(items):
            row = d_table.rows[row_i]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Inches(1.6)
            c1.width = Inches(4.9)
            set_cell_background(c0, HEX_LIGHT_BG)
            set_cell_margins(c0, 100, 100, 120, 120)
            set_cell_margins(c1, 100, 100, 120, 120)

            p0 = c0.paragraphs[0]
            r0 = p0.add_run(k)
            format_font(r0, size_pt=9, bold=True, color=COLOR_DARK)

            p1 = c1.paragraphs[0]
            r1 = p1.add_run(str(v))
            format_font(r1, size_pt=9, bold=False, color=COLOR_DARK)
            p1.paragraph_format.line_spacing = 1.2

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    doc.add_page_break()

    # ==========================================
    # 5. 제4장. 솔루션 평가 (VOC) 및 트릿지 지원 요청서
    # ==========================================
    h4 = doc.add_paragraph()
    r_h4 = h4.add_run("제4장. 솔루션 평가(VOC) 및 트릿지(CSM) 개선 요청서")
    format_font(r_h4, size_pt=16, bold=True, color=COLOR_PRIMARY)
    h4.paragraph_format.space_before = Pt(10)
    h4.paragraph_format.space_after = Pt(12)

    p_voc_desc = doc.add_paragraph()
    r_vd = p_voc_desc.add_run(ai_analysis.get("issues_summary", ""))
    format_font(r_vd, size_pt=10, color=COLOR_DARK)
    p_voc_desc.paragraph_format.line_spacing = 1.3
    p_voc_desc.paragraph_format.space_after = Pt(10)

    # 부서별 이슈 및 요청사항 표
    voc_table = doc.add_table(rows=1, cols=3)
    voc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(voc_table, HEX_BORDER)

    voc_headers = ["부서명 / 솔루션", "현장 발생 이슈 및 애로사항", "트릿지(CSM) 공식 개선 요청사항"]
    for i, h in enumerate(voc_headers):
        cell = voc_table.rows[0].cells[i]
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, 120, 120, 120, 120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        format_font(r, size_pt=9.5, bold=True, color=RGBColor(255, 255, 255))

    voc_table.rows[0].cells[0].width = Inches(1.6)
    voc_table.rows[0].cells[1].width = Inches(2.4)
    voc_table.rows[0].cells[2].width = Inches(2.5)

    for idx, rec in enumerate(records):
        issue_text = rec.get("issue_and_request", "-")
        if not issue_text or issue_text == "-":
            continue
        row = voc_table.add_row()
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(1.6)
        c1.width = Inches(2.4)
        c2.width = Inches(2.5)
        set_cell_background(c0, HEX_LIGHT_BG)
        set_cell_margins(c0, 100, 100, 100, 100)
        set_cell_margins(c1, 100, 100, 100, 100)
        set_cell_margins(c2, 100, 100, 100, 100)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(f"{rec.get('department')}\n({rec.get('solution')})")
        format_font(r0, size_pt=9, bold=True, color=COLOR_DARK)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(issue_text)
        format_font(r1, size_pt=8.5, color=COLOR_DARK)

        p2 = c2.paragraphs[0]
        r2 = p2.add_run("전담 CSM 기술 미팅 안건 등록 및 데이터 보정 요청")
        format_font(r2, size_pt=8.5, color=COLOR_SECONDARY)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # ==========================================
    # 6. 제5장. 종합 제언 및 향후 로드맵
    # ==========================================
    h5 = doc.add_paragraph()
    r_h5 = h5.add_run("제5장. 종합 제언 및 향후 활용 로드맵")
    format_font(r_h5, size_pt=16, bold=True, color=COLOR_PRIMARY)
    h5.paragraph_format.space_before = Pt(10)
    h5.paragraph_format.space_after = Pt(12)

    recs = ai_analysis.get("strategic_recommendations", [])
    for idx, rec_text in enumerate(recs):
        p_rec = doc.add_paragraph()
        r_bullet = p_rec.add_run(f"✔ ")
        format_font(r_bullet, size_pt=11, bold=True, color=COLOR_ACCENT)
        r_body = p_rec.add_run(rec_text)
        format_font(r_body, size_pt=10, bold=False, color=COLOR_DARK)
        p_rec.paragraph_format.line_spacing = 1.3
        p_rec.paragraph_format.space_after = Pt(10)

    # 맺음말
    p_end = doc.add_paragraph()
    r_end = p_end.add_run(
        "\n[총평] 트릿지는 단순한 정보 검색 툴을 넘어 당사의 소싱 원가 절감과 글로벌 수출선 다변화에 "
        "직접적인 기여를 입증하고 있습니다. 향후 부서별 교육 강화와 트릿지 CSM과의 정기 기술 피드백 루프를 구축하여 "
        "구독 가치를 극대화할 것을 제언합니다."
    )
    format_font(r_end, size_pt=9.5, bold=True, color=COLOR_PRIMARY)
    p_end.paragraph_format.line_spacing = 1.3

    # 문서 저장
    doc.save(output_path)
    return output_path
